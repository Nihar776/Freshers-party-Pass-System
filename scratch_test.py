import sys
import os
sys.path.append(os.getcwd())
from database import SessionLocal
from schema_v2 import Student, PaymentStatus
import openpyxl
from openpyxl.styles import Alignment
import io

db = SessionLocal()
students = db.query(Student).filter(Student.payment_status != PaymentStatus.NOT_PURCHASED).limit(10).all()

wb = openpyxl.Workbook()
master_ws = wb.active
master_ws.title = "Master Data"

headers = ["Sr No", "SAP ID", "Name", "Year", "Branch", "Food Pref", "Solo/Group", "Paid By", "Amount Paid", "Seller"]
master_ws.append(headers)

groups_dict = {}
for s in students:
    gid = s.group_id if s.group_id else f"solo_{s.id}"
    if gid not in groups_dict:
        groups_dict[gid] = []
    groups_dict[gid].append(s)
    
group_counter = 1
processed_groups = []

for gid, members in groups_dict.items():
    members.sort(key=lambda s: s.is_group_payer, reverse=True)
    is_group = len(members) > 1 or (members[0].pass_type and members[0].pass_type.value == 'group')
    group_label = f"Group-{group_counter}" if is_group else "Solo"
    if is_group:
        group_counter += 1
        
    payer_name = members[0].name
    total_amount = sum(m.amount or 0 for m in members)
    seller = members[0].distributor.full_name if members[0].distributor else "Public Portal"
    
    processed_groups.append({
        "members": members,
        "label": group_label,
        "payer": payer_name,
        "amount": total_amount,
        "seller": seller
    })
    
sr_no = 1
current_row = 2
center_aligned_text = Alignment(vertical='center', horizontal='center')

for pg in processed_groups:
    start_row = current_row
    for m in pg["members"]:
        food_val = (m.food_preference.value if hasattr(m.food_preference, "value") else str(m.food_preference or "")).lower()
        food_pref = "Jain" if food_val == "jain" else "Non-Jain"
        
        row = [
            sr_no,
            m.sap_id,
            m.name,
            m.year or "",
            m.branch or "",
            food_pref,
            pg["label"],
            pg["payer"],
            pg["amount"],
            pg["seller"]
        ]
        master_ws.append(row)
        sr_no += 1
        current_row += 1
        
    end_row = current_row - 1
    if end_row > start_row:
        master_ws.merge_cells(start_row=start_row, start_column=7, end_row=end_row, end_column=7)
        master_ws.merge_cells(start_row=start_row, start_column=8, end_row=end_row, end_column=8)
        master_ws.merge_cells(start_row=start_row, start_column=9, end_row=end_row, end_column=9)
        for col in (7, 8, 9):
            master_ws.cell(row=start_row, column=col).alignment = center_aligned_text

wb.save('test_sales.xlsx')
print("Done")

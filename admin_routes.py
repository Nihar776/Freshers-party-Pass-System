"""
admin_routes.py
The admin's full-picture dashboard: every stat requested, pulled from the
students/expenses/handovers tables. Also exposes the audit log for review
and a one-click integrity check on the hash chain.
"""
from typing import Optional
import csv
import io
import base64
import openpyxl

from fastapi import APIRouter, Depends, Query, HTTPException, BackgroundTasks, File, Form, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import func
from sqlalchemy.orm import Session

from datetime import datetime, date, timedelta
from database import get_db
from schema_v2 import (
    Student, User, UserRole, PaymentStatus, PaymentMode, PassType,
    FoodPreference, Expense, CashHandover, AuditLog, DiscountCode, DiscountType
)
from session_auth import require_role
from audit import verify_chain_integrity, write_audit_log
from auth import generate_pass_token, generate_qr_image
from mailer import send_pass_email
from distributor_routes import calculate_discount, PASS_PRICE

router = APIRouter(prefix="/admin", tags=["admin-dashboard"])


# ---------------------------------------------------------------------------
# Schemas
# ---------------------------------------------------------------------------
class DistributorSales(BaseModel):
    distributor_name: str
    passes_sold: int
    amount_collected: float


class BranchSales(BaseModel):
    branch: Optional[str]
    passes_sold: int
    percent_of_total: float
    cash_amount: float = 0.0
    cash_percent: float = 0.0
    upi_amount: float = 0.0
    upi_percent: float = 0.0


class PaymentModeBreakdown(BaseModel):
    mode: str
    count: int
    amount: float
    percent_of_count: float
    percent_of_amount: float


class PassTypeBreakdown(BaseModel):
    pass_type: str
    count: int
    percent: float

class FoodPreferenceBreakdown(BaseModel):
    food_preference: str
    count: int
    percent: float

class DashboardResponse(BaseModel):
    # Sales overview
    total_roster_size: int
    total_sold: int             # verified + pending
    total_verified: int
    total_pending_verification: int
    total_rejected: int

    distributor_wise: list[DistributorSales]
    branch_wise: list[BranchSales]
    payment_mode_breakdown: list[PaymentModeBreakdown]
    pass_type_breakdown: list[PassTypeBreakdown]
    food_preference_breakdown: list[FoodPreferenceBreakdown]

    # Money (VERIFIED sales only - pending UPI never counts toward "collected")
    net_collection: float
    net_cash_collection: float
    net_upi_collection: float

    # Expenses / funds
    total_expenses: float
    net_funds_remaining: float
    total_cash_outstanding_with_distributors: float

    # Gate night
    total_entered: int
    pending_entry: int

    total_meals_taken: int
    meals_jain_taken: int
    meals_non_jain_taken: int


class AuditLogEntry(BaseModel):
    id: int
    timestamp: str
    user_name: Optional[str]
    action: str
    table_name: str
    record_id: int
    old_value: Optional[dict]
    new_value: Optional[dict]

    class Config:
        from_attributes = True


class DiscountCodeCreate(BaseModel):
    code: str
    discount_type: DiscountType
    discount_value: float
    max_uses: Optional[int] = None


class DiscountCodeResponse(BaseModel):
    id: int
    code: str
    discount_type: DiscountType
    discount_value: float
    max_uses: Optional[int]
    times_used: int
    is_active: bool
    created_at: str

    class Config:
        from_attributes = True


class StudentOverrideRequest(BaseModel):
    payment_status: Optional[PaymentStatus] = None
    is_used: Optional[bool] = None
    food_preference: Optional[FoodPreference] = None


class StudentBulkUpdateRequest(BaseModel):
    student_ids: list[int]
    action: str
    value: str

class GroupEditRequest(BaseModel):
    sap_ids: list[str]


class AdminStudentSummary(BaseModel):
    id: int
    sap_id: str
    name: str
    email: Optional[str] = None
    branch: str
    payment_status: PaymentStatus
    payment_mode: Optional[PaymentMode] = None
    food_preference: FoodPreference
    is_used: bool
    sold_at: Optional[str]
    group_id: Optional[str] = None
    amount: Optional[float] = None
    serial_number: Optional[str] = None


class AdminSellerDetail(BaseModel):
    id: int
    name: str
    passes_sold: int
    cash_collected: float
    upi_collected: float
    outstanding_cash: float


class AuditIntegrityResult(BaseModel):
    intact: bool
    first_broken_entry_id: Optional[int]
    total_entries: int


ADMIN_ONLY = (UserRole.ADMIN,)


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------
@router.get("/dashboard", response_model=DashboardResponse)
def dashboard(
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY)),
):
    total_roster = db.query(func.count(Student.id)).scalar() or 0

    status_counts = dict(
        db.query(Student.payment_status, func.count(Student.id))
        .group_by(Student.payment_status).all()
    )
    total_verified = status_counts.get(PaymentStatus.VERIFIED, 0)
    total_pending = status_counts.get(PaymentStatus.PENDING_VERIFICATION, 0)
    total_rejected = status_counts.get(PaymentStatus.REJECTED, 0)
    total_sold = total_verified + total_pending

    # --- Distributor-wise (verified sales only - that's real, counted money) ---
    dist_rows = (
        db.query(User.full_name, func.count(Student.id), func.coalesce(func.sum(Student.amount), 0.0))
        .join(Student, Student.distributor_id == User.id)
        .filter(Student.payment_status == PaymentStatus.VERIFIED)
        .group_by(User.full_name)
        .order_by(func.count(Student.id).desc())
        .all()
    )
    distributor_wise = [
        DistributorSales(distributor_name=name, passes_sold=count, amount_collected=amount)
        for name, count, amount in dist_rows
    ]

    # --- Branch-wise ---
    from sqlalchemy import case
    branch_rows = (
        db.query(
            Student.branch,
            func.count(Student.id),
            func.coalesce(func.sum(case((Student.payment_mode == PaymentMode.CASH, Student.amount), else_=0)), 0.0),
            func.coalesce(func.sum(case((Student.payment_mode == PaymentMode.UPI, Student.amount), else_=0)), 0.0)
        )
        .filter(Student.payment_status == PaymentStatus.VERIFIED)
        .group_by(Student.branch)
        .order_by(func.count(Student.id).desc())
        .all()
    )
    branch_wise = [
        BranchSales(
            branch=b, passes_sold=c, 
            percent_of_total=round(100 * c / total_verified, 1) if total_verified else 0.0,
            cash_amount=cash_amt,
            cash_percent=round(100 * cash_amt / (cash_amt + upi_amt), 1) if (cash_amt + upi_amt) > 0 else 0.0,
            upi_amount=upi_amt,
            upi_percent=round(100 * upi_amt / (cash_amt + upi_amt), 1) if (cash_amt + upi_amt) > 0 else 0.0,
        )
        for b, c, cash_amt, upi_amt in branch_rows
    ]

    # --- Payment mode breakdown ---
    mode_rows = (
        db.query(Student.payment_mode, func.count(Student.id), func.coalesce(func.sum(Student.amount), 0.0))
        .filter(Student.payment_status == PaymentStatus.VERIFIED)
        .group_by(Student.payment_mode)
        .all()
    )
    total_amount = sum(amt for _, _, amt in mode_rows) or 1.0  # avoid /0
    payment_mode_breakdown = [
        PaymentModeBreakdown(
            mode=mode.value if mode else "unknown", count=count, amount=amount,
            percent_of_count=round(100 * count / total_verified, 1) if total_verified else 0.0,
            percent_of_amount=round(100 * amount / total_amount, 1),
        )
        for mode, count, amount in mode_rows
    ]
    net_cash = next((r.amount for r in payment_mode_breakdown if r.mode == "cash"), 0.0)
    net_upi = next((r.amount for r in payment_mode_breakdown if r.mode == "upi"), 0.0)

    # --- Pass type breakdown ---
    type_rows = (
        db.query(Student.pass_type, func.count(Student.id))
        .filter(Student.payment_status == PaymentStatus.VERIFIED)
        .group_by(Student.pass_type)
        .all()
    )
    pass_type_breakdown = [
        PassTypeBreakdown(
            pass_type=pt.value, count=c,
            percent=round(100 * c / total_verified, 1) if total_verified else 0.0,
        )
        for pt, c in type_rows
    ]
    
    # --- Food preference breakdown ---
    food_rows = (
        db.query(Student.food_preference, func.count(Student.id))
        .filter(Student.payment_status == PaymentStatus.VERIFIED)
        .group_by(Student.food_preference)
        .all()
    )
    food_breakdown = [
        FoodPreferenceBreakdown(
            food_preference=fp.value if fp else "unknown", count=c,
            percent=round(100 * c / total_verified, 1) if total_verified else 0.0,
        )
        for fp, c in food_rows
    ]

    # --- Cash outstanding across ALL distributors ---
    all_cash_collected = db.query(func.coalesce(func.sum(Student.amount), 0.0)).filter(
        Student.payment_mode == PaymentMode.CASH, Student.payment_status == PaymentStatus.VERIFIED
    ).scalar()
    all_handed_over = db.query(func.coalesce(func.sum(CashHandover.amount), 0.0)).scalar()
    cash_outstanding = all_cash_collected - all_handed_over

    # --- Expenses / net funds ---
    total_expenses = db.query(func.coalesce(func.sum(Expense.amount), 0.0)).scalar()
    net_collection = net_cash + net_upi

    # --- Gate night ---
    total_entered = db.query(func.count(Student.id)).filter(Student.is_used.is_(True)).scalar() or 0

    # --- Food night ---
    meals_taken_rows = (
        db.query(Student.food_preference, func.count(Student.id))
        .filter(Student.food_received.is_(True))
        .group_by(Student.food_preference)
        .all()
    )
    meals_jain = 0
    meals_veg = 0
    for fp, count in meals_taken_rows:
        if fp and fp.value == "jain":
            meals_jain += count
        elif fp and fp.value == "veg":
            meals_veg += count
    
    total_meals = meals_jain + meals_veg

    return DashboardResponse(
        total_roster_size=total_roster,
        total_sold=total_sold,
        total_verified=total_verified,
        total_pending_verification=total_pending,
        total_rejected=total_rejected,
        distributor_wise=distributor_wise,
        branch_wise=branch_wise,
        payment_mode_breakdown=payment_mode_breakdown,
        pass_type_breakdown=pass_type_breakdown,
        food_preference_breakdown=food_breakdown,
        net_collection=net_collection,
        net_cash_collection=net_cash,
        net_upi_collection=net_upi,
        total_expenses=total_expenses,
        net_funds_remaining=net_collection - total_expenses,
        total_cash_outstanding_with_distributors=cash_outstanding,
        total_entered=total_entered,
        pending_entry=total_verified - total_entered,
        total_meals_taken=total_meals,
        meals_jain_taken=meals_jain,
        meals_non_jain_taken=meals_veg,
    )


@router.get("/students", response_model=list[AdminStudentSummary])
def list_students_admin(
    search: Optional[str] = None,
    status: Optional[str] = None,
    limit: int = 100,
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY)),
):
    q = db.query(Student)
    if search:
        q = q.filter((Student.name.ilike(f"%{search}%")) | (Student.sap_id.ilike(f"%{search}%")))
    if status == "entered":
        q = q.filter(Student.is_used == True)
    elif status == "bought":
        q = q.filter(Student.payment_status == PaymentStatus.VERIFIED)
    elif status == "pending":
        q = q.filter(Student.payment_status == PaymentStatus.PENDING_VERIFICATION)
    elif status == "not_purchased":
        q = q.filter(Student.payment_status == PaymentStatus.NOT_PURCHASED)

    # Calculate global serial numbers for ALL sold passes
    all_sold = db.query(Student.id, Student.group_id).filter(
        Student.payment_status.in_([PaymentStatus.VERIFIED, PaymentStatus.PENDING_VERIFICATION])
    ).order_by(Student.sold_at).all()
    
    group_counter = 1
    group_map = {}
    serial_map = {}
    
    for s_id, g_id in all_sold:
        if not g_id:
            serial_map[s_id] = str(group_counter)
            group_counter += 1
        else:
            if g_id not in group_map:
                group_map[g_id] = {"num": group_counter, "idx": 1}
                group_counter += 1
            g_info = group_map[g_id]
            serial_map[s_id] = f"{g_info['num']}.{g_info['idx']}"
            g_info["idx"] += 1

    rows = q.order_by(Student.name).limit(limit).all()
    return [
        AdminStudentSummary(
            id=s.id,
            sap_id=s.sap_id,
            name=s.name,
            email=s.email,
            branch=s.branch,
            payment_status=s.payment_status,
            payment_mode=s.payment_mode,
            food_preference=s.food_preference,
            is_used=s.is_used,
            sold_at=s.sold_at.isoformat() if s.sold_at else None,
            group_id=s.group_id,
            amount=s.amount,
            serial_number=serial_map.get(s.id)
        )
        for s in rows
    ]


@router.get("/students/{student_id}/screenshot")
def get_student_screenshot(
    student_id: int,
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY)),
):
    student = db.query(Student).filter(Student.id == student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")
        
    if not student.payment_screenshot:
        raise HTTPException(status_code=404, detail="No screenshot uploaded for this student")
        
    return {"screenshot": student.payment_screenshot}


@router.get("/sellers-detail", response_model=list[AdminSellerDetail])
def sellers_detail(
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY)),
):
    sellers = db.query(User).filter(User.role.in_([UserRole.DISTRIBUTOR, UserRole.ADMIN])).all()
    result = []
    for s in sellers:
        # Passes sold (verified)
        passes_sold = db.query(func.count(Student.id)).filter(
            Student.distributor_id == s.id, Student.payment_status == PaymentStatus.VERIFIED
        ).scalar() or 0

        # Cash collected
        cash_collected = db.query(func.coalesce(func.sum(Student.amount), 0.0)).filter(
            Student.distributor_id == s.id, 
            Student.payment_status == PaymentStatus.VERIFIED,
            Student.payment_mode == PaymentMode.CASH
        ).scalar()
        
        # UPI collected
        upi_collected = db.query(func.coalesce(func.sum(Student.amount), 0.0)).filter(
            Student.distributor_id == s.id, 
            Student.payment_status == PaymentStatus.VERIFIED,
            Student.payment_mode == PaymentMode.UPI
        ).scalar()

        # Handed over
        handed_over = db.query(func.coalesce(func.sum(CashHandover.amount), 0.0)).filter(
            CashHandover.distributor_id == s.id
        ).scalar()

        result.append(AdminSellerDetail(
            id=s.id, name=s.full_name, passes_sold=passes_sold,
            cash_collected=cash_collected, upi_collected=upi_collected,
            outstanding_cash=cash_collected - handed_over
        ))
    return result


# ---------------------------------------------------------------------------
# Audit log review
# ---------------------------------------------------------------------------
@router.get("/audit-log", response_model=list[AuditLogEntry])
def view_audit_log(
    action: Optional[str] = None,
    limit: int = Query(100, le=500),
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY)),
):
    q = db.query(AuditLog)
    if action:
        q = q.filter(AuditLog.action == action)
    rows = q.order_by(AuditLog.id.desc()).limit(limit).all()
    return [
        AuditLogEntry(
            id=r.id, timestamp=r.timestamp.isoformat(),
            user_name=r.user.full_name if r.user else None,
            action=r.action, table_name=r.table_name, record_id=r.record_id,
            old_value=r.old_value, new_value=r.new_value,
        )
        for r in rows
    ]


@router.get("/audit-log/verify-integrity", response_model=AuditIntegrityResult)
def check_audit_integrity(
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY)),
):
    intact, bad_id = verify_chain_integrity(db)
    total = db.query(func.count(AuditLog.id)).scalar() or 0
    return AuditIntegrityResult(intact=intact, first_broken_entry_id=bad_id, total_entries=total)


# ---------------------------------------------------------------------------
# Export & Overrides
# ---------------------------------------------------------------------------
@router.get("/export/attendees")
def export_attendees(
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY)),
):
    students = db.query(Student).filter(Student.is_used.is_(True)).all()
    
    wb = openpyxl.Workbook()
    master_ws = wb.active
    master_ws.title = "Master Data"
    
    headers = ["Name", "SAP ID", "Branch", "Entered At", "Scanned By", "Food Pref", "Food Taken", "Food Taken At", "Food Scanned By"]
    master_ws.append(headers)
    
    branch_sheets = {}
    
    for s in students:
        scanned_by = s.scanned_by.full_name if s.scanned_by else "Unknown"
        entered_at = s.entered_at.strftime("%Y-%m-%d %H:%M:%S") if s.entered_at else "Unknown"
        food_pref = "Jain" if (s.food_preference and s.food_preference.value == "jain") else "Non-Jain"
        food_taken = "Yes" if s.food_received else "No"
        food_taken_at = s.food_received_at.strftime("%Y-%m-%d %H:%M:%S") if s.food_received_at else ""
        food_scanned = s.food_scanned_by.full_name if s.food_scanned_by else ""
        row = [s.name, s.sap_id, s.branch, entered_at, scanned_by, food_pref, food_taken, food_taken_at, food_scanned]
        
        master_ws.append(row)
        
        branch_name = s.branch or "Unknown"
        if branch_name not in branch_sheets:
            # Excel limits sheet name to 31 chars and bans some special chars
            safe_title = "".join(c for c in branch_name if c.isalnum() or c in " _-")[:31]
            if not safe_title: safe_title = "Unknown"
            # Ensure unique
            while safe_title in wb.sheetnames:
                safe_title = safe_title[:28] + "_1"
            branch_sheets[branch_name] = wb.create_sheet(title=safe_title)
            branch_sheets[branch_name].append(headers)
            
        branch_sheets[branch_name].append(row)
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output, 
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 
        headers={"Content-Disposition": "attachment; filename=attendees.xlsx"}
    )


@router.get("/export/sales")
def export_sales(
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY)),
):
    students = db.query(Student).filter(Student.payment_status != PaymentStatus.NOT_PURCHASED).all()
    
    wb = openpyxl.Workbook()
    master_ws = wb.active
    master_ws.title = "Master Data"
    
    headers = ["Name", "SAP ID", "Branch", "Pass Type", "Amount", "Payment Mode", "Status", "Distributor"]
    master_ws.append(headers)
    
    branch_sheets = {}
    
    for s in students:
        distributor = s.distributor.full_name if s.distributor else "Unknown"
        pass_type = s.pass_type.value if s.pass_type else ""
        payment_mode = s.payment_mode.value if s.payment_mode else ""
        row = [s.name, s.sap_id, s.branch, pass_type, s.amount or 0, payment_mode, s.payment_status.value, distributor]
        
        master_ws.append(row)
        
        branch_name = s.branch or "Unknown"
        if branch_name not in branch_sheets:
            safe_title = "".join(c for c in branch_name if c.isalnum() or c in " _-")[:31]
            if not safe_title: safe_title = "Unknown"
            while safe_title in wb.sheetnames:
                safe_title = safe_title[:28] + "_1"
            branch_sheets[branch_name] = wb.create_sheet(title=safe_title)
            branch_sheets[branch_name].append(headers)
            
        branch_sheets[branch_name].append(row)
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output, 
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 
        headers={"Content-Disposition": "attachment; filename=sales.xlsx"}
    )


@router.patch("/students/{student_id}/override")
async def override_student(
    student_id: int,
    payment_status: Optional[PaymentStatus] = Form(None),
    is_used: Optional[bool] = Form(None),
    food_preference: Optional[FoodPreference] = Form(None),
    payment_mode: Optional[PaymentMode] = Form(None),
    utr_number: Optional[str] = Form(None),
    screenshot: Optional[UploadFile] = File(None),
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY)),
):
    student = db.query(Student).filter(Student.id == student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")

    old_snapshot = {
        "payment_status": student.payment_status.value if student.payment_status else None,
        "is_used": student.is_used,
        "food_preference": student.food_preference.value if student.food_preference else None,
        "payment_mode": student.payment_mode.value if student.payment_mode else None,
        "utr_number": student.utr_number
    }
    
    new_snapshot = {}
    if payment_status is not None:
        student.payment_status = payment_status
        new_snapshot["payment_status"] = payment_status.value
    if is_used is not None:
        student.is_used = is_used
        new_snapshot["is_used"] = is_used
    if food_preference is not None:
        student.food_preference = food_preference
        new_snapshot["food_preference"] = food_preference.value
    if payment_mode is not None:
        student.payment_mode = payment_mode
        new_snapshot["payment_mode"] = payment_mode.value
    if utr_number is not None:
        student.utr_number = utr_number
        new_snapshot["utr_number"] = utr_number
        
    if screenshot is not None:
        ss_bytes = await screenshot.read()
        if ss_bytes:
            student.payment_screenshot = base64.b64encode(ss_bytes).decode("utf-8")
            new_snapshot["payment_screenshot"] = "Uploaded"

    if not new_snapshot:
        return {"message": "No changes requested"}

    db.flush()

    write_audit_log(
        db, user_id=admin.id, action="admin_override", table_name="students",
        record_id=student.id, old_value=old_snapshot, new_value=new_snapshot,
    )
    db.commit()
    db.refresh(student)

    return {"message": "Student record overridden successfully"}


import csv
import io

@router.post("/upload-food-csv")
async def upload_food_csv(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY))
):
    if not file.filename.endswith(".csv"):
        raise HTTPException(status_code=400, detail="Must be a CSV file")

    content = await file.read()
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        raise HTTPException(status_code=400, detail="File is not valid UTF-8")

    reader = csv.DictReader(io.StringIO(text))
    
    # Normalize headers
    if not reader.fieldnames:
        raise HTTPException(status_code=400, detail="CSV is empty or has no headers")
    headers = [h.strip().lower() for h in reader.fieldnames]
    reader.fieldnames = headers

    if "sap_id" not in headers or "food_preference" not in headers:
        raise HTTPException(status_code=400, detail="CSV must contain 'sap_id' and 'food_preference' columns")

    updated = 0
    errors = []

    for idx, row in enumerate(reader, start=2):
        sap_id = row.get("sap_id", "").strip()
        pref_raw = row.get("food_preference", "").strip().lower()

        if not sap_id or not pref_raw:
            continue

        if pref_raw not in ("veg", "jain"):
            errors.append(f"Row {idx}: Invalid preference '{pref_raw}' for SAP ID {sap_id}")
            continue

        student = db.query(Student).filter(Student.sap_id == sap_id).first()
        if not student:
            errors.append(f"Row {idx}: Student with SAP ID {sap_id} not found")
            continue

        old_val = student.food_preference.value if student.food_preference else None
        if old_val != pref_raw:
            student.food_preference = FoodPreference(pref_raw)
            write_audit_log(
                db, user_id=admin.id, action="bulk_food_update", table_name="students",
                record_id=student.id, old_value={"food_preference": old_val}, new_value={"food_preference": pref_raw}
            )
            updated += 1

    db.commit()
    
    msg = f"Updated {updated} records."
    if errors:
        msg += f" {len(errors)} errors occurred (check logs)."
        print("Bulk food update errors:", errors)
        
    return {"message": msg}


@router.patch("/groups/{group_id}")
def edit_group(
    group_id: str,
    payload: GroupEditRequest,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY))
):
    if not payload.sap_ids:
        raise HTTPException(status_code=400, detail="A group must have at least one member.")

    # Fetch all students currently in the group
    current_members = db.query(Student).filter(Student.group_id == group_id).all()
    if not current_members:
        raise HTTPException(status_code=404, detail="Group not found.")
        
    current_sap_ids = {s.sap_id for s in current_members}
    target_sap_ids = set(payload.sap_ids)
    
    # Identify to_remove and to_add
    saps_to_remove = current_sap_ids - target_sap_ids
    saps_to_add = target_sap_ids - current_sap_ids
    
    # 1. Process Removals
    if saps_to_remove:
        members_to_remove = [s for s in current_members if s.sap_id in saps_to_remove]
        for m in members_to_remove:
            old_snapshot = {
                "group_id": m.group_id,
                "payment_status": m.payment_status.value if m.payment_status else None,
                "amount": m.amount,
                "is_group_payer": m.is_group_payer
            }
            m.group_id = None
            m.payment_status = PaymentStatus.NOT_PURCHASED
            m.amount = None
            m.is_group_payer = False
            write_audit_log(
                db, user_id=admin.id, action="group_member_removed", table_name="students",
                record_id=m.id, old_value=old_snapshot, new_value={"group_id": None, "payment_status": PaymentStatus.NOT_PURCHASED.value}
            )

    # 2. Process Additions
    added_students = []
    if saps_to_add:
        students_to_add = db.query(Student).filter(Student.sap_id.in_(saps_to_add)).all()
        found_saps = {s.sap_id for s in students_to_add}
        missing = saps_to_add - found_saps
        if missing:
            raise HTTPException(status_code=404, detail=f"SAP IDs not found: {', '.join(missing)}")
            
        # Ensure they are available
        for m in students_to_add:
            if m.payment_status not in (PaymentStatus.NOT_PURCHASED, PaymentStatus.REJECTED):
                raise HTTPException(status_code=400, detail=f"Student {m.sap_id} is already in a group or purchased.")
            if not m.email:
                raise HTTPException(status_code=400, detail=f"Student {m.sap_id} has no email. Update their email first.")
                
            old_snapshot = {
                "group_id": m.group_id,
                "payment_status": m.payment_status.value if m.payment_status else None
            }
            m.group_id = group_id
            m.payment_status = current_members[0].payment_status # Match the group's status
            m.sold_at = current_members[0].sold_at
            m.payment_mode = current_members[0].payment_mode
            m.distributor_id = current_members[0].distributor_id
            
            write_audit_log(
                db, user_id=admin.id, action="group_member_added", table_name="students",
                record_id=m.id, old_value=old_snapshot, new_value={"group_id": group_id, "payment_status": m.payment_status.value}
            )
            added_students.append(m)
            
    # 3. Recalculate Group Amount
    # Fetch final group members
    final_members = db.query(Student).filter(Student.group_id == group_id).all()
    group_size = len(final_members)
    
    if group_size > 0:
        base_sold_at = final_members[0].sold_at
        group_discount_pct = calculate_discount(group_size, base_sold_at)
        total_base_amount = PASS_PRICE * group_size
        
        # We don't retroactively apply a new Promo code if they didn't have one, 
        # but if they DID have one, we need to respect it. This requires checking discount_code_id.
        dc_id = final_members[0].discount_code_id
        dc = db.query(DiscountCode).filter(DiscountCode.id == dc_id).first() if dc_id else None
        
        if dc:
            if dc.discount_type == DiscountType.PERCENTAGE:
                code_discount_pct = dc.discount_value / 100.0
                best_discount_pct = max(group_discount_pct, code_discount_pct)
                final_total_amount = total_base_amount * (1.0 - best_discount_pct)
            elif dc.discount_type == DiscountType.FLAT:
                pct_discount_amount = total_base_amount * group_discount_pct
                best_discount_amount = max(pct_discount_amount, dc.discount_value)
                final_total_amount = max(0.0, total_base_amount - best_discount_amount)
        else:
            final_total_amount = total_base_amount * (1.0 - group_discount_pct)
            
        amount_per_student = final_total_amount / group_size
        
        for m in final_members:
            if m.amount != amount_per_student:
                old_amount = m.amount
                m.amount = amount_per_student
                write_audit_log(
                    db, user_id=admin.id, action="group_price_recalculated", table_name="students",
                    record_id=m.id, old_value={"amount": old_amount}, new_value={"amount": amount_per_student}
                )
                
    db.commit()
    
    # 4. Trigger emails for any newly added members if the group is verified
    if added_students and final_members and final_members[0].payment_status == PaymentStatus.VERIFIED:
        for m in added_students:
            token = generate_pass_token(sap_id=m.sap_id, pass_uuid=m.pass_uuid)
            qr_image = generate_qr_image(token)
            background_tasks.add_task(
                send_pass_email,
                recipient_email=m.email,
                student_name=m.name,
                qr_image_bytes=qr_image,
                sap_id=m.sap_id,
            )

    return {"message": f"Group updated. Size is now {group_size}."}


@router.post("/students/bulk-update")
def bulk_update_students(
    payload: StudentBulkUpdateRequest,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY)),
):
    if not payload.student_ids:
        raise HTTPException(status_code=400, detail="No students selected")

    students = db.query(Student).filter(Student.id.in_(payload.student_ids)).all()
    if not students:
        raise HTTPException(status_code=404, detail="No matching students found")

    updated_count = 0
    for student in students:
        old_snapshot = {
            "payment_status": student.payment_status.value if student.payment_status else None,
            "is_used": student.is_used,
            "food_preference": student.food_preference.value if student.food_preference else None
        }
        new_snapshot = {}

        if payload.action == "payment_status":
            try:
                new_status = PaymentStatus(payload.value)
                if student.payment_status != new_status:
                    student.payment_status = new_status
                    new_snapshot["payment_status"] = new_status.value
            except ValueError:
                continue
        elif payload.action == "is_used":
            new_used = payload.value.lower() == "true"
            if student.is_used != new_used:
                student.is_used = new_used
                new_snapshot["is_used"] = new_used
        elif payload.action == "food_preference":
            try:
                new_food = FoodPreference(payload.value)
                if student.food_preference != new_food:
                    student.food_preference = new_food
                    new_snapshot["food_preference"] = new_food.value
            except ValueError:
                continue
        elif payload.action == "resend_emails":
            if student.payment_status != PaymentStatus.VERIFIED or not student.email:
                continue
            token = generate_pass_token(sap_id=student.sap_id, pass_uuid=student.pass_uuid)
            qr_image = generate_qr_image(token)
            background_tasks.add_task(
                send_pass_email,
                recipient_email=student.email,
                student_name=student.name,
                qr_image_bytes=qr_image,
                sap_id=student.sap_id,
            )
            updated_count += 1
            continue

        if new_snapshot:
            db.flush()
            write_audit_log(
                db, user_id=admin.id, action="admin_override", table_name="students",
                record_id=student.id, old_value=old_snapshot, new_value=new_snapshot,
            )
            updated_count += 1

    db.commit()
    return {"message": f"Successfully updated {updated_count} students"}


# ---------------------------------------------------------------------------
# Discount Codes
# ---------------------------------------------------------------------------
@router.post("/discount-codes", response_model=DiscountCodeResponse)
def create_discount_code(
    payload: DiscountCodeCreate,
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY)),
):
    code_upper = payload.code.upper().strip()
    existing = db.query(DiscountCode).filter(func.upper(DiscountCode.code) == code_upper).first()
    if existing:
        raise HTTPException(status_code=400, detail="Discount code already exists")

    if payload.discount_value <= 0:
        raise HTTPException(status_code=400, detail="Discount value must be greater than 0")

    if payload.discount_type == DiscountType.PERCENTAGE and payload.discount_value > 100:
        raise HTTPException(status_code=400, detail="Percentage discount cannot exceed 100")

    dc = DiscountCode(
        code=code_upper,
        discount_type=payload.discount_type,
        discount_value=payload.discount_value,
        max_uses=payload.max_uses,
        created_by_id=admin.id
    )
    db.add(dc)
    db.flush()

    write_audit_log(
        db, user_id=admin.id, action="discount_code_created",
        table_name="discount_codes", record_id=dc.id,
        new_value={"code": dc.code, "type": dc.discount_type.value, "value": dc.discount_value, "max": dc.max_uses}
    )
    db.commit()
    db.refresh(dc)

    return DiscountCodeResponse(
        id=dc.id, code=dc.code, discount_type=dc.discount_type,
        discount_value=dc.discount_value, max_uses=dc.max_uses,
        times_used=dc.times_used, is_active=dc.is_active,
        created_at=dc.created_at.isoformat()
    )


@router.post("/resend-email/{sap_id}")
def resend_email(
    sap_id: str,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    user: User = Depends(require_role(UserRole.ADMIN, UserRole.DISTRIBUTOR)),
):
    student = db.query(Student).filter(Student.sap_id == sap_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")
    if student.payment_status != PaymentStatus.VERIFIED:
        raise HTTPException(status_code=400, detail="Cannot resend email - pass is not verified")
    if not student.email:
        raise HTTPException(status_code=400, detail="No email address on file for this student")

    token = generate_pass_token(sap_id=student.sap_id, pass_uuid=student.pass_uuid)
    qr_image = generate_qr_image(token)
    background_tasks.add_task(
        send_pass_email,
        recipient_email=student.email,
        student_name=student.name,
        qr_image_bytes=qr_image,
        sap_id=student.sap_id,
    )

    return {"message": "Email is being sent in the background"}


@router.get("/discount-codes", response_model=list[DiscountCodeResponse])
def list_discount_codes(
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY)),
):
    codes = db.query(DiscountCode).order_by(DiscountCode.created_at.desc()).all()
    return [
        DiscountCodeResponse(
            id=c.id, code=c.code, discount_type=c.discount_type,
            discount_value=c.discount_value, max_uses=c.max_uses,
            times_used=c.times_used, is_active=c.is_active,
            created_at=c.created_at.isoformat()
        ) for c in codes
    ]


@router.patch("/discount-codes/{code_id}/toggle", response_model=DiscountCodeResponse)
def toggle_discount_code(
    code_id: int,
    db: Session = Depends(get_db),
    admin: User = Depends(require_role(*ADMIN_ONLY)),
):
    dc = db.query(DiscountCode).filter(DiscountCode.id == code_id).first()
    if not dc:
        raise HTTPException(status_code=404, detail="Discount code not found")

    old_active = dc.is_active
    dc.is_active = not dc.is_active

    db.flush()
    write_audit_log(
        db, user_id=admin.id, action="discount_code_toggled", table_name="discount_codes",
        record_id=dc.id, old_value={"is_active": old_active}, new_value={"is_active": dc.is_active}
    )
    db.commit()
    db.refresh(dc)

    return DiscountCodeResponse(
        id=dc.id, code=dc.code, discount_type=dc.discount_type,
        discount_value=dc.discount_value, max_uses=dc.max_uses,
        times_used=dc.times_used, is_active=dc.is_active,
        created_at=dc.created_at.isoformat()
    )


# ---------------------------------------------------------------------------
# Global Settings
# ---------------------------------------------------------------------------
from settings_manager import get_settings, save_settings

class SettingsPayload(BaseModel):
    scanner_mode: str

@router.get("/settings")
def get_global_settings(admin: User = Depends(require_role(*ADMIN_ONLY))):
    return get_settings()

@router.post("/settings")
def update_global_settings(payload: SettingsPayload, admin: User = Depends(require_role(*ADMIN_ONLY))):
    settings = get_settings()
    settings["scanner_mode"] = payload.scanner_mode
    save_settings(settings)
    return {"message": "Settings updated"}
